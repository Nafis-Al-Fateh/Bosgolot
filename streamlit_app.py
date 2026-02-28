import streamlit as st
import fitz
import pytesseract
import cv2
import numpy as np
import tempfile
import os
from docx import Document
from openpyxl import Workbook
from PIL import Image
import shutil

st.set_page_config(page_title="Bangla PDF OCR System", layout="wide")
st.title("📄 PDF → DOCX + Excel (Image-Based Multi-Page OCR)")
st.markdown("High Resolution Image OCR | Bangla + English | No Encoding Errors")

# -----------------------------
# IMAGE PREPROCESSING
# -----------------------------
def preprocess_image(image):

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)

    # contrast enhancement
    gray = cv2.equalizeHist(gray)

    # noise removal
    gray = cv2.fastNlMeansDenoising(gray, None, 30, 7, 21)

    # adaptive threshold
    thresh = cv2.adaptiveThreshold(
        gray,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY,
        15,
        8
    )

    return thresh


# -----------------------------
# OCR WITH STRUCTURED DATA
# -----------------------------
def ocr_structured(image):

    data = pytesseract.image_to_data(
        image,
        lang="ben+eng",
        config="--oem 3 --psm 6",
        output_type=pytesseract.Output.DICT
    )

    words = []
    n = len(data["text"])

    for i in range(n):
        if int(data["conf"][i]) > 40:
            text = data["text"][i].strip()
            if text:
                words.append({
                    "text": text,
                    "x": data["left"][i],
                    "y": data["top"][i]
                })

    return words


# -----------------------------
# GROUP WORDS INTO LINES
# -----------------------------
def group_into_lines(words):

    lines = {}
    threshold = 15  # vertical grouping threshold

    for word in words:
        y = word["y"]

        found = False
        for key in lines.keys():
            if abs(key - y) < threshold:
                lines[key].append(word)
                found = True
                break

        if not found:
            lines[y] = [word]

    sorted_lines = []

    for key in sorted(lines.keys()):
        line_words = sorted(lines[key], key=lambda x: x["x"])
        line_text = " ".join([w["text"] for w in line_words])
        sorted_lines.append(line_text)

    return sorted_lines


# -----------------------------
# PROCESS PDF
# -----------------------------
def process_pdf(pdf_path):

    pdf = fitz.open(pdf_path)

    document = Document()
    workbook = Workbook()
    workbook.remove(workbook.active)

    for page_index in range(len(pdf)):

        page = pdf[page_index]

        # Convert page to high resolution image
        pix = page.get_pixmap(dpi=400)
        img = np.frombuffer(pix.samples, dtype=np.uint8)
        img = img.reshape(pix.height, pix.width, pix.n)

        if pix.n == 4:
            img = cv2.cvtColor(img, cv2.COLOR_BGRA2BGR)

        processed = preprocess_image(img)

        # OCR structured
        words = ocr_structured(processed)

        # Group into lines
        lines = group_into_lines(words)

        # Add to DOCX
        for line in lines:
            document.add_paragraph(line)

        document.add_page_break()

        # Add to Excel (each page = new sheet)
        sheet = workbook.create_sheet(f"Page_{page_index+1}")

        for line in lines:
            sheet.append([line])

    # Save DOCX
    doc_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    document.save(doc_file.name)

    # Save Excel
    excel_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    workbook.save(excel_file.name)

    return doc_file.name, excel_file.name


# -----------------------------
# STREAMLIT UI
# -----------------------------
uploaded_file = st.file_uploader("Upload PDF (Max 10MB)", type=["pdf"])

if uploaded_file:

    if uploaded_file.size > 10 * 1024 * 1024:
        st.error("File exceeds 10MB limit.")
    else:

        temp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        temp_pdf.write(uploaded_file.read())
        temp_pdf.close()

        if st.button("Process PDF"):

            try:
                with st.spinner("Processing multi-page OCR..."):

                    doc_path, excel_path = process_pdf(temp_pdf.name)

                st.success("Processing Complete!")

                with open(doc_path, "rb") as f:
                    st.download_button(
                        "Download DOCX",
                        f,
                        file_name="output.docx"
                    )

                with open(excel_path, "rb") as f:
                    st.download_button(
                        "Download Excel",
                        f,
                        file_name="output.xlsx"
                    )

                os.remove(temp_pdf.name)
                os.remove(doc_path)
                os.remove(excel_path)

            except Exception as e:
                st.error("Processing failed. Please restart and try again.")
                st.exception(e)
